"""recurring advisor workflow regression tests."""
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace as NS
from unittest.mock import patch
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.core.database import SessionLocal
from app.models import Holding, PerformanceSnapshot, Investor, FollowUpTask, AdviceRecord, SourceRow, ImportBatch, User
from app.engines.patterns import comparable_date
from tests.conftest import sign_in
from tests.test_phase2 import make_workbook, upload, supervisor_login


def headers(csrf):return {'Origin':'http://localhost:5173','X-CSRF-Token':csrf}

def commit(client,token,content):
    staged=upload(client,token,content)
    assert staged.status_code==201,staged.text
    return client.post('/api/v1/imports/'+staged.json()['id']+'/commit',headers=headers(token))

def later_workbook(conflict=False):
    wb=load_workbook(BytesIO(make_workbook()))
    wb['holdings']['C2']='2026-08-31'
    wb['holdings']['J2']=12000
    wb['performance_history']['B2']='2026-08-01'
    wb['performance_history']['C2']=12000
    wb['investors']['B2']='Updated source label'
    if conflict:wb['holdings']['C2']='2026-07-31'
    out=BytesIO();wb.save(out);return out.getvalue()

def test_later_import_keeps_history_and_skips_old_transactions(client):
    token=supervisor_login(client)['csrf_token']
    assert commit(client,token,make_workbook()).status_code==200
    later=later_workbook()
    second=commit(client,token,later)
    assert second.status_code==200,second.text
    assert second.json()['unchanged_observations_skipped']==1
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Holding))==2
        assert db.scalar(select(func.count()).select_from(PerformanceSnapshot))==2
        assert db.scalar(select(func.count()).select_from(Investor).where(Investor.id=='DEMO-001'))==1
    assert upload(client,token,later).status_code==409

def test_conflicting_correction_rolls_back_whole_import(client):
    token=supervisor_login(client)['csrf_token']
    assert commit(client,token,make_workbook()).status_code==200
    assert commit(client,token,later_workbook(conflict=True)).status_code==409
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Holding))==1
        assert db.scalar(select(func.count()).select_from(PerformanceSnapshot))==1

def test_manual_review_and_meeting_task_need_no_detected_pattern(client):
    h=headers(sign_in(client)['csrf_token'])
    case=client.post('/api/v1/attention',headers=h,json={'investor_id':'A-001','title':'Discuss future liquidity','reason':'Client requested a review of next year funding needs.'})
    assert case.status_code==201,case.text
    assert client.post('/api/v1/attention',headers=h,json={'investor_id':'A-002','title':'Unauthorized review','reason':'This should not expose another advisor client.'}).status_code==404
    note=client.post('/api/v1/meetings/notes',headers=h,json={'investor_id':'A-001','transcript':'Advisor to request the latest statement from the client.'}).json()
    title='Request the latest statement from the client'
    payload={'source_meeting_id':note['id'],'commitment_index':0,'title':title,'due_date':(date.today()+timedelta(days=3)).isoformat()}
    assert client.post('/api/v1/tasks',headers=h,json=payload).status_code==422
    assert client.post(f"/api/v1/meetings/{note['id']}/approve",headers=h,json={'summary':'Agreed to obtain an updated client investment statement.','commitments':[title]}).status_code==200
    task=client.post('/api/v1/tasks',headers=h,json=payload)
    assert task.status_code==201,task.text
    assert task.json()['advice_id'] is None
    assert task.json()['source_meeting_id']==note['id']
    sign_in(client,'b@advicue.local')
    assert client.get('/api/v1/tasks').json()['items']==[]

def test_full_transcript_extraction_includes_tail_and_is_atomic(client):
    h=headers(sign_in(client)['csrf_token'])
    transcript='Discussion about liquidity. '*600+'TAIL_COMMITMENT: obtain statement.'
    note=client.post('/api/v1/meetings/notes',headers=h,json={'investor_id':'A-001','transcript':transcript}).json()
    chunks=[]
    def model(system,facts, **kwargs):
        chunks.append(facts['transcript'])
        return {'summary':'Reviewed documented client liquidity needs.','commitments':['Obtain statement from client'] if 'TAIL_COMMITMENT' in facts['transcript'] else [],'review_questions':[]}
    with patch('app.api.local_ai.local_llm_json',side_effect=model):
        result=client.post(f"/api/v1/meetings/{note['id']}/extract",headers=h)
    assert result.status_code==200,result.text
    draft=result.json()['extracted_draft']
    assert draft['coverage']['characters_processed']==len(transcript)
    assert draft['coverage']['complete']
    assert draft['commitments']==['Obtain statement from client']
    assert any('TAIL_COMMITMENT' in c for c in chunks)

def test_same_date_coverage_survives_one_newer_client():
    rows=[NS(as_of_date=date(2026,7,31),investor_id=str(i)) for i in range(8)]
    rows.append(NS(as_of_date=date(2026,8,31),investor_id='0'))
    assert comparable_date(rows)==date(2026,7,31)

def test_goal_plan_saved_and_scoped(client):
    with SessionLocal() as db:
        inv=db.get(Investor,'A-001');inv.goal_target_amount=1000000;inv.goal_target_date=date.today()+timedelta(days=1500);db.commit()
    h=headers(sign_in(client)['csrf_token'])
    payload={'principal':'50000','monthly_contribution':'1000','as_of_date':date.today().isoformat(),'annual_return_assumptions_pct':[3,6,9]}
    result=client.put('/api/v1/investors/A-001/goal-plan',headers=h,json=payload)
    assert result.status_code==200,result.text
    assert client.get('/api/v1/investors/A-001/goal-plan').json()['plan']['monthly_contribution']=='1000.00'
    payload['annual_return_assumptions_pct']=[-100]
    assert client.put('/api/v1/investors/A-001/goal-plan',headers=h,json=payload).status_code==422
    sign_in(client,'b@advicue.local')
    assert client.get('/api/v1/investors/A-001/goal-plan').status_code==404

def test_measurable_advice_links_to_later_holdings(client):
    h=headers(sign_in(client)['csrf_token'])
    case=client.post('/api/v1/attention',headers=h,json={'investor_id':'A-001','title':'Review agreed allocation','reason':'Confirm the documented allocation threshold at review.'}).json()
    due=date.today().isoformat()
    advice=client.post('/api/v1/advice',headers=h,json={'alert_id':case['id'],'recommendation':'Review the agreed client allocation condition.','rationale':'A measurable observation should be checked after approval.','outcome_rule':{'metric':'equity_pct','comparator':'>=','target_pct':20,'due_date':due}})
    assert advice.status_code==201,advice.text
    assert client.post(f"/api/v1/advice/{advice.json()['id']}/finalize",headers=h,json={'note':'Advisor confirms this recorded condition for review.'}).status_code==200
    with SessionLocal() as db:
        obj=db.get(AdviceRecord,advice.json()['id']);obj.decided_at=datetime.now(timezone.utc)-timedelta(days=10)
        batch=ImportBatch(organization_id='org-a',uploaded_by='advisor-a',filename='later.xlsx',file_sha256='v4later');db.add(batch);db.flush()
        for idx,(cls,value) in enumerate([('Direct Equity',300),('Fixed Deposit',700)]):
            row=SourceRow(batch_id=batch.id,sheet_name='holdings',row_number=idx+2,raw_data={},normalized_data={});db.add(row);db.flush()
            db.add(Holding(organization_id='org-a',investor_id='A-001',source_row_id=row.id,external_holding_id='later'+str(idx),as_of_date=date.today(),asset_class=cls,instrument_name=cls,current_value=value))
        db.commit()
    result=client.get(f"/api/v1/advice/{advice.json()['id']}/outcome").json()
    assert result['condition_check']['status']=='SATISFIED_OBSERVATION'
    assert result['condition_check']['deadline_verified'] is True
    assert len(result['condition_check']['evidence_source_row_ids'])==2
    assert result['task_completion']['done']==0
    assert result['status']=='NOT_ASSESSABLE' # no investment return evidence

def test_complete_advisor_journey_with_second_import(client,monkeypatch):
    from tests.test_advisor_only import demo_permissions
    demo_permissions(monkeypatch)
    with SessionLocal() as db:
        db.get(User,'advisor-a').external_advisor_code='ADV05';db.commit()
    token=sign_in(client)['csrf_token'];h=headers(token)
    def workbook(day,equity=False):
        wb=load_workbook(BytesIO(make_workbook()))
        wb['holdings']['C2']=day.isoformat();wb['performance_history']['B2']=day.isoformat()
        if equity:
            wb['holdings']['D2']='Direct Equity';wb['holdings']['E2']='Example Equity';wb['holdings']['J2']=12000
            wb['performance_history']['C2']=12000
        out=BytesIO();wb.save(out);return out.getvalue()
    assert commit(client,token,workbook(date.today()-timedelta(days=20))).status_code==200
    case=client.post('/api/v1/attention',headers=h,json={'investor_id':'DEMO-001','title':'Review the agreed equity condition','reason':'Client requested a measured follow-up on the recorded allocation.'}).json()
    advice=client.post('/api/v1/advice',headers=h,json={'alert_id':case['id'],'recommendation':'Check the allocation condition agreed with the client.','rationale':'Use dated evidence after the advisor review to assess the condition.','outcome_rule':{'metric':'equity_pct','comparator':'>=','target_pct':20,'due_date':date.today().isoformat()}})
    assert advice.status_code==201,advice.text
    with patch('app.api.workflow._now',return_value=datetime.now(timezone.utc)-timedelta(days=10)):
        finalized=client.post(f"/api/v1/advice/{advice.json()['id']}/finalize",headers=h,json={'note':'Advisor records the client-agreed measurable condition.'})
    assert finalized.status_code==200,finalized.text
    before=client.get(f"/api/v1/advice/{advice.json()['id']}/outcome").json()
    assert before['condition_check']['status']=='NOT_ASSESSABLE'
    assert commit(client,token,workbook(date.today(),True)).status_code==200
    after=client.get(f"/api/v1/advice/{advice.json()['id']}/outcome").json()
    assert after['condition_check']['status']=='SATISFIED_OBSERVATION'
    assert after['condition_check']['observed_pct']=='100.00'
    assert after['condition_check']['deadline_verified']
    assert after['status']=='OBSERVATIONAL_ONLY'
    with SessionLocal() as db:assert db.scalar(select(func.count()).select_from(Holding))==2
