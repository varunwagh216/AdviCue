from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from sqlalchemy import func, select
from app.core.database import SessionLocal
from app.models import (AuditEvent, DataQualityIssue, Holding, ImportBatch, InvestmentTransaction,
                        Investor, PerformanceSnapshot, SourceRow, User)
from tests.conftest import sign_in

# Tests generate data fixtures so the distributed source project does not contain personal records.
DATA_FILE = Path('/mnt/data/Transaction History Investor.xlsx')


def make_workbook(bad=False, duplicate=False):
    wb=Workbook()
    wb.remove(wb.active)
    sheets={
        'investors': ["investor_id","name","age","advisor_id","risk_profile","primary_goal","goal_target_amount","goal_target_date","onboarding_date","city","effective_date"],
        'holdings': ["investor_id","holding_id","as_of_date","asset_class","instrument_name","units","avg_purchase_price","purchase_date","current_price","current_value"],
        'transactions': ["investor_id","transaction_id","date","type","instrument_name","asset_class","units","price_per_unit","amount"],
        'performance_history': ["investor_id","snapshot_date","portfolio_value","benchmark_index_value","portfolio_return_pct","benchmark_return_pct"],
    }
    for sn,heads in sheets.items():
        ws=wb.create_sheet(sn); ws.append(heads)
    wb['investors'].append(['DEMO-001','Demo Investor',35,'ADV05','Moderate','Retirement','₹10,00,000','2035-01-01','01-02-2025','Pune','2025-02-01'])
    wb['holdings'].append(['DEMO-001','H001','2026-07-31','Fixed Deposit','Bank FD',100,100,'01-03-2025',110,11000])
    wb['transactions'].append(['DEMO-001','TX001','2025-04-03','SIP','Fund','Fixed Deposit',1,100,100])
    if duplicate:
        wb['transactions'].append(['DEMO-001','TX002','2025-04-03','SIP','Fund','Fixed Deposit',1,100,100])
    wb['performance_history'].append(['DEMO-001','2026-07-01',11000,1000,1.2,1.1])
    if bad:
        wb['holdings'].append(['MISSING','H002','2026-07-31','Fixed Deposit','FD 2',2,100,'2025-02-01',100,-200])
    b=BytesIO(); wb.save(b); return b.getvalue()


def supervisor_login(client):
    return sign_in(client,'supervisor@advicue.local')


def upload(client, csrf, content, filename='investor.xlsx'):
    return client.post('/api/v1/imports', files={'file':(filename,content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
                       headers={'X-CSRF-Token':csrf,'Origin':'http://localhost:5173'})


def test_supervisor_upload_review_commit(client):
    auth=supervisor_login(client)
    original_workbook=make_workbook(duplicate=True)
    r=upload(client,auth['csrf_token'],original_workbook)
    assert r.status_code==201, r.text
    body=r.json()
    assert body['status']=='STAGED'
    assert body['sheet_counts']=={'investors':1,'holdings':1,'transactions':2,'performance_history':1}
    assert body['total_rows']==5 and body['error_count']==0 and body['warning_count']>=1
    issues=client.get(f"/api/v1/imports/{body['id']}/issues").json()
    assert any(x['code']=='DUPLICATE_CANDIDATE' for x in issues['items'])
    assert client.get('/api/v1/imports').json()['items'][0]['id']==body['id']
    assert client.post(f"/api/v1/imports/{body['id']}/commit").status_code==403  # CSRF
    result=client.post(f"/api/v1/imports/{body['id']}/commit",headers={'X-CSRF-Token':auth['csrf_token'],'Origin':'http://localhost:5173'})
    assert result.status_code==200,result.text
    assert result.json()['status']=='COMMITTED'
    ids={x['id'] for x in client.get('/api/v1/investors').json()['items']}
    assert ids=={'A-001','A-002','DEMO-001'}
    details=client.get('/api/v1/investors/DEMO-001')
    assert details.status_code==200
    assert len(details.json()['holdings'])==1
    assert details.json()['transaction_count']==2
    assert details.json()['performance_status']=='NOT_ASSESSABLE'
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(SourceRow))==5
        assert db.scalar(select(func.count()).select_from(Holding))==1
        assert db.scalar(select(func.count()).select_from(InvestmentTransaction))==2
        assert db.scalar(select(func.count()).select_from(PerformanceSnapshot))==1
        assert db.scalar(select(func.count()).select_from(InvestmentTransaction).where(InvestmentTransaction.possible_duplicate.is_(True)))==1
        inv=db.get(Investor,'DEMO-001')
        assert inv.source_advisor_code=='ADV05' and inv.goal_target_amount==1000000
        assert db.get(User,inv.advisor_id).is_active is False
        assert db.scalar(select(func.count()).select_from(AuditEvent).where(AuditEvent.action=='IMPORT_COMMITTED'))==1
    again=upload(client,auth['csrf_token'],original_workbook)
    assert again.status_code==409 and body['id'] in again.json()['detail']
    again_commit=client.post(f"/api/v1/imports/{body['id']}/commit", headers={'X-CSRF-Token':auth['csrf_token']})
    assert again_commit.status_code==409


def test_bad_records_block_commit_but_preserve_raw(client):
    auth=supervisor_login(client)
    up=upload(client,auth['csrf_token'],make_workbook(bad=True))
    assert up.status_code==201,up.text
    body=up.json()
    assert body['status']=='BLOCKED' and body['error_count']>0
    errs=client.get(f"/api/v1/imports/{body['id']}/issues").json()['items']
    assert any(x['code']=='INVESTOR_NOT_FOUND' for x in errs)
    assert any(x['code']=='NEGATIVE_VALUE' for x in errs)
    commit=client.post(f"/api/v1/imports/{body['id']}/commit",headers={'X-CSRF-Token':auth['csrf_token']})
    assert commit.status_code==409
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(SourceRow))==5
        assert db.scalar(select(func.count()).select_from(Holding))==0
        bad_row=db.scalar(select(SourceRow).where(SourceRow.sheet_name=='holdings',SourceRow.row_number==3))
        assert bad_row.raw_data['current_value']==-200
        assert bad_row.accepted is False


def test_permissions_origin_and_file_size(client):
    assert client.get('/api/v1/imports').status_code==401
    advisor=sign_in(client)
    assert client.get('/api/v1/imports').status_code==403
    assert upload(client,advisor['csrf_token'],make_workbook()).status_code==403
    client.post('/api/v1/auth/logout',headers={'X-CSRF-Token':advisor['csrf_token']})
    su=supervisor_login(client)
    assert upload(client,su['csrf_token'],make_workbook(),filename='evil.csv').status_code==422
    assert upload(client,su['csrf_token'],b'junk').status_code==422
    assert client.post('/api/v1/imports',files={'file':('investor.xlsx',make_workbook())},headers={'X-CSRF-Token':su['csrf_token'],'Origin':'https://evil.test'}).status_code==403
    assert client.post('/api/v1/imports',files={'file':('investor.xlsx',make_workbook())}).status_code==403
    assert upload(client,su['csrf_token'],b'a'*(10*1024*1024+1)).status_code==413


def test_duplicate_existing_investor_blocks_atomic_commit(client):
    # A-001 already exists in Phase 1 fixtures; ensure zero normalized records are added.
    wb=Workbook();wb.remove(wb.active)
    from app.services.importer import SHEETS
    for sh,columns in SHEETS.items():wb.create_sheet(sh).append(columns)
    wb['investors'].append(['A-001','Existing',30,'ADV02','Moderate',None,None,None,None,None,None])
    b=BytesIO();wb.save(b)
    su=supervisor_login(client)
    staged=upload(client,su['csrf_token'],b.getvalue())
    assert staged.status_code==201,staged.text
    batch_id=staged.json()['id']
    result=client.post(f'/api/v1/imports/{batch_id}/commit',headers={'X-CSRF-Token':su['csrf_token']})
    assert result.status_code==409
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Investor))==3
        assert db.scalar(select(func.count()).select_from(SourceRow))==1
        assert db.get(ImportBatch,batch_id).status=='STAGED'


def test_organization_isolation(client):
    su=supervisor_login(client)
    up=upload(client,su['csrf_token'],make_workbook())
    assert up.status_code==201
    batch_id=up.json()['id']
    client.post('/api/v1/auth/logout',headers={'X-CSRF-Token':su['csrf_token']})
    sign_in(client,'c@advicue.local')
    assert client.get(f'/api/v1/imports/{batch_id}').status_code==403


def test_actual_uploaded_workbook_if_present(client):
    if not DATA_FILE.exists():
        return
    su=supervisor_login(client)
    up=upload(client,su['csrf_token'],DATA_FILE.read_bytes(),DATA_FILE.name)
    assert up.status_code==201,up.text
    data=up.json()
    assert data['sheet_counts']=={'investors':48,'holdings':208,'transactions':394,'performance_history':576}
    assert data['error_count']==0, data
    committed=client.post(f"/api/v1/imports/{data['id']}/commit",headers={'X-CSRF-Token':su['csrf_token']})
    assert committed.status_code==200,committed.text
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Investor))==51
        assert db.scalar(select(func.count()).select_from(Holding))==208
        assert db.scalar(select(func.count()).select_from(InvestmentTransaction))==394
        assert db.scalar(select(func.count()).select_from(PerformanceSnapshot))==576
        assert len({i.source_advisor_code for i in db.scalars(select(Investor).where(Investor.id.like('INV%'))).all()})==8


def test_mapped_advisor_sees_only_their_imported_client(client):
    with SessionLocal() as db:
        u=db.get(User,'advisor-a')
        u.external_advisor_code='ADV05'
        db.commit()
    supervisor=supervisor_login(client)
    up=upload(client,supervisor['csrf_token'],make_workbook())
    assert up.status_code==201
    bid=up.json()['id']
    assert client.post(f'/api/v1/imports/{bid}/commit',headers={'X-CSRF-Token':supervisor['csrf_token']}).status_code==200
    client.post('/api/v1/auth/logout',headers={'X-CSRF-Token':supervisor['csrf_token']})
    sign_in(client)
    ids={x['id'] for x in client.get('/api/v1/investors').json()['items']}
    assert ids=={'A-001','DEMO-001'}
    assert client.get('/api/v1/investors/DEMO-001').status_code==200
    assert client.get('/api/v1/investors/A-002').status_code==404
    assert client.get('/api/v1/investors/B-001').status_code==404


def test_import_isolation_between_supervisors(client):
    supervisor=supervisor_login(client)
    up=upload(client,supervisor['csrf_token'],make_workbook())
    assert up.status_code==201
    bid=up.json()['id']
    client.post('/api/v1/auth/logout',headers={'X-CSRF-Token':supervisor['csrf_token']})
    other=sign_in(client,'supervisor-c@advicue.local')
    assert client.get(f'/api/v1/imports/{bid}').status_code==404
    assert client.get(f'/api/v1/imports/{bid}/issues').status_code==404
    assert client.post(f'/api/v1/imports/{bid}/commit',headers={'X-CSRF-Token':other['csrf_token']}).status_code==404
