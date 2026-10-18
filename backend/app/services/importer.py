"""Excel import: preserve raw evidence, validate deterministically, then commit atomically.

NOT a performance-return engine. Snapshot % columns remain source-reported / unverified.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from zipfile import BadZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.services.trust import sip_recurrence_days, date_classification
from app.models import (AuditEvent, DataQualityIssue, Holding, ImportBatch,
                        InvestmentTransaction, Investor, PerformanceSnapshot,
                        SourceRow, User)

SHEETS = {
    'investors': ['investor_id','name','age','advisor_id','risk_profile','primary_goal','goal_target_amount','goal_target_date','onboarding_date','city','effective_date'],
    'holdings': ['investor_id','holding_id','as_of_date','asset_class','instrument_name','units','avg_purchase_price','purchase_date','current_price','current_value'],
    'transactions': ['investor_id','transaction_id','date','type','instrument_name','asset_class','units','price_per_unit','amount'],
    'performance_history': ['investor_id','snapshot_date','portfolio_value','benchmark_index_value','portfolio_return_pct','benchmark_return_pct'],
}
COLUMNS_DATES = {'goal_target_date','onboarding_date','effective_date','as_of_date','purchase_date','date','snapshot_date'}
COLUMNS_MONEY = {'goal_target_amount','units','avg_purchase_price','current_price','current_value','price_per_unit','amount','portfolio_value','benchmark_index_value','portfolio_return_pct','benchmark_return_pct'}
REQUIRED_TEXT = {
    'investors': ['investor_id','name','advisor_id'],
    'holdings': ['investor_id','holding_id','as_of_date','asset_class','instrument_name','current_value'],
    'transactions': ['investor_id','transaction_id','date','type','instrument_name','asset_class','amount'],
    'performance_history': ['investor_id','snapshot_date','portfolio_value'],
}
MAX_BYTES = 10 * 1024 * 1024
MAX_DATA_ROWS = 200_000


def jsonable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (bool, str, int, float)) or value is None:
        return value
    return str(value)


def text(value: Any) -> str | None:
    if value is None or str(value).strip() == '':
        return None
    return str(value).strip()


def parse_date(value: Any) -> str | None:
    if value is None or text(value) is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', s):
        return date.fromisoformat(s).isoformat()
    if re.fullmatch(r'\d{2}-\d{2}-\d{4}', s):
        return datetime.strptime(s,'%d-%m-%Y').date().isoformat()
    if re.fullmatch(r'\d{2}/\d{2}/\d{4}', s):
        return datetime.strptime(s,'%d/%m/%Y').date().isoformat()
    raise ValueError('Expected YYYY-MM-DD, DD-MM-YYYY or DD/MM/YYYY')


def parse_number(value: Any) -> str | None:
    if value is None or text(value) is None:
        return None
    if isinstance(value, bool):
        raise ValueError('Boolean is not a numeric value')
    s = str(value).strip().replace(',','').replace('₹','')
    s = re.sub(r'^(INR|Rs\.?)\s*', '', s, flags=re.IGNORECASE).strip()
    try:
        d = Decimal(s)
    except (InvalidOperation, ValueError):
        raise ValueError('Invalid number') from None
    if not d.is_finite():
        raise ValueError('Number must be finite')
    if len(d.as_tuple().digits) > 22 or abs(d.as_tuple().exponent) > 12:
        raise ValueError('Number exceeds supported precision')
    return str(d)


def issue(issues: list[dict], row: SourceRow, severity: str, code: str, message: str, field: str | None = None):
    issues.append({'source_row_id':row.id,'sheet_name':row.sheet_name,'row_number':row.row_number,
                   'field_name':field,'severity':severity,'code':code,'message':message})
    if severity == 'ERROR':
        row.accepted = False


def normalize(row: SourceRow, issues: list[dict]):
    result = {}
    for key, value in row.raw_data.items():
        try:
            if key in COLUMNS_DATES:
                result[key] = parse_date(value)
                # Preserve the decision; never silently call ambiguous day-first text exact.
                raw=str(value).strip() if value is not None else ''
                if re.fullmatch(r'(\d{2})[-/](\d{2})[-/](\d{4})',raw):
                    day,month=map(int,raw[:5].replace('/','-').split('-'))
                    label='AMBIGUOUS_DAY_FIRST' if day<=12 and month<=12 else 'EXPLICIT_DAY_FIRST'
                    result[f'{key}_date_quality']=label
                    if label=='AMBIGUOUS_DAY_FIRST':
                        issue(issues,row,'WARNING','AMBIGUOUS_SOURCE_DATE',
                              f'{key} can be read as either day-first or month-first; provisional day-first interpretation, verify before cash-flow calculations.',key)
                elif result[key]:result[f'{key}_date_quality']='EXACT_ISO_OR_EXCEL'
            elif key in COLUMNS_MONEY:
                result[key] = parse_number(value)
            elif key == 'age':
                if value is None or str(value).strip() == '':
                    result[key] = None
                else:
                    d = Decimal(str(value))
                    if d != d.to_integral_value():
                        raise ValueError('Age must be a whole number')
                    result[key] = int(d)
            else:
                result[key] = text(value)
        except (ValueError, InvalidOperation, OverflowError) as e:
            result[key] = None
            issue(issues, row, 'ERROR','INVALID_VALUE',f'{key}: {e}',key)
    for key in REQUIRED_TEXT[row.sheet_name]:
        if result.get(key) is None:
            issue(issues, row, 'ERROR','MISSING_REQUIRED',f'{key} is required',key)
    if row.sheet_name == 'investors':
        age = result.get('age')
        if age is not None and not 18 <= age <= 120:
            issue(issues,row,'WARNING','AGE_OUTSIDE_RANGE','Investor age is outside 18–120','age')
        amount = result.get('goal_target_amount')
        if amount is not None and Decimal(amount) < 0:
            issue(issues,row,'ERROR','NEGATIVE_GOAL','Goal target cannot be negative','goal_target_amount')
    elif row.sheet_name == 'holdings':
        for key in ('units','avg_purchase_price','current_price','current_value'):
            if result.get(key) is not None and Decimal(result[key]) < 0:
                issue(issues,row,'ERROR','NEGATIVE_VALUE',f'{key} cannot be negative',key)
        if all(result.get(k) is not None for k in ('units','current_price','current_value')):
            expected = Decimal(result['units']) * Decimal(result['current_price'])
            actual = Decimal(result['current_value'])
            if abs(expected-actual) > max(Decimal('2'),abs(actual)*Decimal('0.005')):
                issue(issues,row,'WARNING','VALUATION_MISMATCH',
                      'Units × price differs from recorded value by more than the review tolerance','current_value')
    elif row.sheet_name in ('transactions','performance_history'):
        for key in (('units','price_per_unit','amount') if row.sheet_name=='transactions' else ('portfolio_value','benchmark_index_value')):
            if result.get(key) is not None and Decimal(result[key]) < 0:
                issue(issues,row,'ERROR','NEGATIVE_VALUE',f'{key} cannot be negative',key)
    if row.sheet_name=='transactions' and result.get('type') and result.get('amount'):
        if result['type'].strip().lower()=='sip' and all(result.get(x) is not None for x in ('units','price_per_unit')):
            expected=Decimal(result['units'])*Decimal(result['price_per_unit'])
            amount=Decimal(result['amount'])
            if abs(expected-amount)>max(Decimal('2'),amount*Decimal('.005')):
                result['_unit_price_quality']='UNRELIABLE_AMOUNT_AUTHORITATIVE'
                issue(issues,row,'WARNING','UNRELIABLE_SIP_UNITS',
                      'Recorded SIP units x price do not match cash amount; use recorded amount only, not units or price, for cash-flow summaries.', 'units')
        elif all(result.get(x) is not None for x in ('units','price_per_unit')):
            result['_unit_price_quality']='MATCHED' if abs(Decimal(result['units'])*Decimal(result['price_per_unit'])-Decimal(result['amount']))<=max(Decimal('2'),Decimal(result['amount'])*Decimal('.005')) else 'MISMATCH'
    row.normalized_data = result


def read_rows(content: bytes) -> tuple[dict[str,list[tuple[int,dict]]],dict[str,int]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, ValueError, OSError, KeyError) as e:
        raise HTTPException(status_code=422, detail='Invalid Excel workbook (.xlsx)') from e
    try:
        missing = [s for s in SHEETS if s not in wb.sheetnames]
        if missing:
            raise HTTPException(status_code=422,detail=f'Missing required sheets: {", ".join(missing)}')
        sheets: dict[str,list[tuple[int,dict]]] = {}
        counts: dict[str,int] = {}
        total = 0
        for sheet_name, expected_headers in SHEETS.items():
            ws = wb[sheet_name]
            rows = ws.iter_rows(values_only=True)
            first = next(rows, None)
            if not first:
                raise HTTPException(status_code=422,detail=f'{sheet_name} has no header row')
            heads = [str(h).strip() if h is not None else '' for h in first]
            if heads[:len(expected_headers)] != expected_headers:
                raise HTTPException(status_code=422,detail=f'{sheet_name} headers do not match the expected workbook schema')
            record_list = []
            for rn, values in enumerate(rows,start=2):
                if all(v is None or str(v).strip()=='' for v in values):
                    continue
                total += 1
                if total > MAX_DATA_ROWS:
                    raise HTTPException(status_code=413,detail='Workbook contains too many data rows')
                record_list.append((rn, {key:jsonable(values[i] if i < len(values) else None) for i,key in enumerate(expected_headers)}))
            sheets[sheet_name] = record_list
            counts[sheet_name] = len(record_list)
        return sheets,counts
    finally:
        wb.close()


def stage_workbook(db: Session, user: User, filename: str, content: bytes) -> ImportBatch:
    if not filename.lower().endswith('.xlsx') or len(content)>MAX_BYTES:
        raise HTTPException(status_code=413 if len(content)>MAX_BYTES else 422,
                            detail='Choose an .xlsx file of 10 MB or less')
    if not content:
        raise HTTPException(status_code=422,detail='Empty upload')
    sha=hashlib.sha256(content).hexdigest()
    existing=db.scalar(select(ImportBatch).where(ImportBatch.organization_id==user.organization_id,ImportBatch.file_sha256==sha))
    if existing:
        raise HTTPException(status_code=409,detail=f'Workbook already uploaded in batch {existing.id} ({existing.status}). Open the existing batch.')
    sheets,counts=read_rows(content)
    batch=ImportBatch(organization_id=user.organization_id, uploaded_by=user.id, filename=filename[:240],
                      file_sha256=sha, sheet_counts=counts, total_rows=sum(counts.values()))
    db.add(batch)
    db.flush()
    issues=[]
    data_rows:dict[str,list[SourceRow]]={s:[] for s in SHEETS}
    for sheet_name, items in sheets.items():
        for rn,raw in items:
            row=SourceRow(batch_id=batch.id,sheet_name=sheet_name,row_number=rn,raw_data=raw,normalized_data={},accepted=True)
            # Defaults on mapped primary keys are generated at flush; explicitly assign for pre-flush audit references.
            from uuid import uuid4
            row.id=str(uuid4())
            normalize(row,issues)
            db.add(row)
            data_rows[sheet_name].append(row)

    # Reclassify provisional day-first dates using independent ISO SIP recurrence.
    expected_days=sip_recurrence_days(data_rows['transactions'])
    for row in data_rows['transactions']:
        d=row.normalized_data
        d['date_date_quality']=date_classification(row.raw_data.get('date'),
                       investor_id=d.get('investor_id'),expected_days=expected_days.get(d.get('investor_id')))
        row.normalized_data=d
    investor_ids={r.normalized_data.get('investor_id') for r in data_rows['investors'] if r.accepted}
    seen_investors=set()
    seen_holding_ids=set()
    for r in data_rows['investors']:
        investor_id=r.normalized_data.get('investor_id')
        if investor_id in seen_investors:
            issue(issues,r,'ERROR','DUPLICATE_ID','Duplicate investor_id in workbook','investor_id')
        if investor_id:
            seen_investors.add(investor_id)
    for sheet_name in ('holdings','transactions','performance_history'):
        for r in data_rows[sheet_name]:
            inv=r.normalized_data.get('investor_id')
            if inv not in investor_ids:
                issue(issues,r,'ERROR','INVESTOR_NOT_FOUND','Investor not found in valid investor rows','investor_id')
            if sheet_name=='holdings':
                h=r.normalized_data.get('holding_id')
                if h in seen_holding_ids:
                    issue(issues,r,'WARNING','DUPLICATE_HOLDING_ID','Holding ID is repeated; verify before relying on it','holding_id')
                if h: seen_holding_ids.add(h)
    fingerprints={}
    for r in data_rows['transactions']:
        normalized=r.normalized_data
        fprint=json.dumps({k:v for k,v in normalized.items() if k!='transaction_id'},sort_keys=True)
        if fprint in fingerprints:
            issue(issues,r,'WARNING','DUPLICATE_CANDIDATE',
                  f'Possible duplicate of transaction sheet row {fingerprints[fprint]}; both rows are preserved')
        else:
            fingerprints[fprint]=r.row_number
    seen_snapshots=set()
    for r in data_rows['performance_history']:
        key=(r.normalized_data.get('investor_id'),r.normalized_data.get('snapshot_date'))
        if key in seen_snapshots:
            issue(issues,r,'WARNING','DUPLICATE_SNAPSHOT','Repeated investor/date snapshot; both rows are preserved')
        seen_snapshots.add(key)
    # Validate cross-sheet consistency without fabricating client mandate decisions.
    risk_equities={}
    hold_values={}
    for r in data_rows['holdings']:
        if not r.accepted:continue
        d=r.normalized_data; key=d.get('investor_id')
        if key and d.get('current_value'):
            hold_values.setdefault(key,{}).setdefault(d.get('asset_class'),Decimal('0'))
            hold_values[key][d.get('asset_class')]+=Decimal(d['current_value'])
    equity_names={'Direct Equity','Equity Mutual Fund','International Equity Fund'}
    from statistics import median
    profile_alloc={}
    for row in data_rows['investors']:
        inv=row.normalized_data.get('investor_id');risk=row.normalized_data.get('risk_profile')
        classes=hold_values.get(inv,{})
        total=sum(classes.values(),Decimal('0'))
        if risk and total>0:
            eq=sum((v for k,v in classes.items() if k in equity_names),Decimal('0'))*100/total
            risk_equities[inv]=eq
            profile_alloc.setdefault(risk.strip().lower(),[]).append(eq)
        amount=row.normalized_data.get('goal_target_amount')
        target=row.normalized_data.get('goal_target_date')
        if amount and target and inv in hold_values:
            try:
                horizon=(date.fromisoformat(target)-date.today()).days
                if 0<horizon<550 and Decimal(amount)>total*8:
                    issue(issues,row,'WARNING','GOAL_PLAUSIBILITY',
                          'Near-term goal exceeds recorded current portfolio value by over 8x; verify planned funding and target.', 'goal_target_amount')
            except (ValueError,TypeError):pass
    medians={k:Decimal(str(median(v))) for k,v in profile_alloc.items()}
    if all(k in medians for k in ('conservative','moderate','aggressive')):
        if not (medians['conservative']<=medians['moderate']<=medians['aggressive']):
            for row in data_rows['investors']:
                issue(issues,row,'WARNING','RISK_LABEL_COHORT_INVERSION',
                      'Observed median equity exposure is not ordered Conservative <= Moderate <= Aggressive; verify risk labels and mandates.', 'risk_profile')
    batch.error_count=sum(x['severity']=='ERROR' for x in issues)
    batch.warning_count=sum(x['severity']=='WARNING' for x in issues)
    if batch.error_count:
        batch.status='BLOCKED'
    db.flush()  # Source rows must exist before FK-linked issues are inserted.
    for x in issues:
        db.add(DataQualityIssue(batch_id=batch.id,**x))
    db.add(AuditEvent(organization_id=user.organization_id,actor_id=user.id,action='IMPORT_STAGED',
                      entity_type='ImportBatch',entity_id=batch.id,
                      detail=json.dumps({'counts':counts,'errors':batch.error_count,'warnings':batch.warning_count})))
    try:
        db.commit()
    except IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=409,detail='Workbook staging conflicted with existing records; no partial data saved') from e
    db.refresh(batch)
    return batch


def commit_workbook(db:Session,user:User,batch:ImportBatch)->dict:
    if batch.status=='COMMITTED':
        raise HTTPException(status_code=409,detail='Batch already committed')
    if batch.status!='STAGED' or batch.error_count:
        raise HTTPException(status_code=409,detail='Batch has blocking issues and cannot be committed')
    rows=db.scalars(select(SourceRow).where(SourceRow.batch_id==batch.id).order_by(SourceRow.row_number)).all()
    by_sheet={s:[] for s in SHEETS}
    for r in rows:
        by_sheet[r.sheet_name].append(r)
    ids=[r.normalized_data['investor_id'] for r in by_sheet['investors']]
    prior={i.id:i for i in db.scalars(select(Investor).where(Investor.id.in_(ids))).all()}
    if any(i.organization_id != user.organization_id for i in prior.values()):
        raise HTTPException(status_code=409,detail='An investor identifier conflicts with another organization. No records changed.')
    for row in by_sheet['investors']:
        d=row.normalized_data
        if d['investor_id'] in prior and prior[d['investor_id']].source_advisor_code != d['advisor_id']:
            raise HTTPException(status_code=409,detail='Advisor reassignment requires an explicit access review; import was not committed.')
    skipped=0
    def duplicate_or_conflict(model, filters, values):
        nonlocal skipped
        db.flush()
        matches=db.scalars(select(model).where(model.organization_id==user.organization_id,*filters)).all()
        if not matches:return False
        def same(row):
            for key,value in values.items():
                stored=getattr(row,key)
                if stored is None or value is None:
                    if stored is not None or value is not None:return False
                elif isinstance(stored,Decimal):
                    if stored!=Decimal(str(value)):return False
                elif stored!=value:return False
            return True
        if any(same(row) for row in matches):
            skipped+=1
            return True
        raise HTTPException(status_code=409,detail='A dated record conflicts with an existing observation. Existing history is preserved; correct the source or use a new observation date.')
    code_set={r.normalized_data['advisor_id'] for r in by_sheet['investors']}
    existing=db.scalars(select(User).where(User.organization_id==user.organization_id,User.external_advisor_code.in_(code_set))).all()
    by_code={u.external_advisor_code:u for u in existing}
    if len(existing) != len(by_code):
        raise HTTPException(status_code=409,detail='Two users share the same external advisor code')
    from uuid import uuid4
    for code in sorted(code_set):
        if code not in by_code:
            u=User(organization_id=user.organization_id,name=f'Imported Advisor {code}',
                   email=f'imported-{uuid4().hex[:16]}@advicue.invalid',role='ADVISOR',is_active=False,
                   password_hash='DISABLED_NO_LOGIN',external_advisor_code=code)
            db.add(u)
            db.flush()
            by_code[code]=u
    for row in by_sheet['investors']:
        d=row.normalized_data
        if d['investor_id'] in prior:
            # Imported profiles remain source-versioned; approved advisor preferences are never overwritten.
            continue
        db.add(Investor(id=d['investor_id'],organization_id=user.organization_id,
                        advisor_id=by_code[d['advisor_id']].id,source_advisor_code=d['advisor_id'],
                        display_name=d['name'][:120],risk_profile=d.get('risk_profile'),age=d.get('age'),
                        primary_goal=d.get('primary_goal'),goal_target_amount=d.get('goal_target_amount'),
                        goal_target_date=date.fromisoformat(d['goal_target_date']) if d.get('goal_target_date') else None,
                        onboarding_date=date.fromisoformat(d['onboarding_date']) if d.get('onboarding_date') else None,
                        effective_date=date.fromisoformat(d['effective_date']) if d.get('effective_date') else None,
                        city=d.get('city'),source_row_id=row.id))
    db.flush()
    duplicate_rows={x.source_row_id for x in db.scalars(select(DataQualityIssue).where(DataQualityIssue.batch_id==batch.id,DataQualityIssue.code=='DUPLICATE_CANDIDATE')).all()}
    for row in by_sheet['holdings']:
        d=row.normalized_data
        if duplicate_or_conflict(Holding,[Holding.investor_id==d['investor_id'],Holding.external_holding_id==d['holding_id'],Holding.as_of_date==date.fromisoformat(d['as_of_date'])],
            {'asset_class':d['asset_class'],'instrument_name':d['instrument_name'],'current_value':d['current_value'],'units':d.get('units'),'current_price':d.get('current_price')}):continue
        db.add(Holding(organization_id=user.organization_id,investor_id=d['investor_id'],source_row_id=row.id,
                       external_holding_id=d['holding_id'],as_of_date=date.fromisoformat(d['as_of_date']),
                       asset_class=d['asset_class'],instrument_name=d['instrument_name'],units=d.get('units'),
                       average_purchase_price=d.get('avg_purchase_price'),
                       purchase_date=date.fromisoformat(d['purchase_date']) if d.get('purchase_date') else None,
                       purchase_date_confidence=d.get('purchase_date_date_quality','SOURCE_UNVERIFIED'),
                       current_price=d.get('current_price'),current_value=d['current_value']))
    for row in by_sheet['transactions']:
        d=row.normalized_data
        if duplicate_or_conflict(InvestmentTransaction,[InvestmentTransaction.investor_id==d['investor_id'],InvestmentTransaction.external_transaction_id==d['transaction_id']],
            {'date':date.fromisoformat(d['date']),'amount':d['amount'],'type':d['type'],'asset_class':d['asset_class'],'instrument_name':d['instrument_name']}):continue
        semantic_copy=db.scalar(select(InvestmentTransaction.id).where(InvestmentTransaction.organization_id==user.organization_id,InvestmentTransaction.investor_id==d['investor_id'],InvestmentTransaction.date==date.fromisoformat(d['date']),InvestmentTransaction.type==d['type'],InvestmentTransaction.amount==Decimal(d['amount']),InvestmentTransaction.instrument_name==d['instrument_name'],InvestmentTransaction.asset_class==d['asset_class']))
        if semantic_copy:
            db.add(DataQualityIssue(batch_id=batch.id,source_row_id=row.id,sheet_name='transactions',row_number=row.row_number,severity='WARNING',code='DUPLICATE_CANDIDATE',message='Same transaction signature exists in imported history; later copy excluded from arithmetic.'))
            batch.warning_count+=1
        db.add(InvestmentTransaction(organization_id=user.organization_id,investor_id=d['investor_id'],source_row_id=row.id,
                                     external_transaction_id=d['transaction_id'],date=date.fromisoformat(d['date']),
                                     type=d['type'],instrument_name=d['instrument_name'],asset_class=d['asset_class'],
                                     units=d.get('units'),price_per_unit=d.get('price_per_unit'),amount=d['amount'],
                                     possible_duplicate=row.id in duplicate_rows or bool(semantic_copy),
                                     date_confidence=d.get('date_date_quality','SOURCE_UNVERIFIED'),
                                     unit_price_confidence=d.get('_unit_price_quality','UNKNOWN'),
                                     external_flow_status= ('CONTRIBUTION_UNVERIFIED' if d['type'].strip().lower()=='sip' else
                                          'WITHDRAWAL_UNVERIFIED' if d['type'].strip().lower()=='redemption' else 'UNKNOWN')))
    for row in by_sheet['performance_history']:
        d=row.normalized_data
        if duplicate_or_conflict(PerformanceSnapshot,[PerformanceSnapshot.investor_id==d['investor_id'],PerformanceSnapshot.snapshot_date==date.fromisoformat(d['snapshot_date'])],
            {'portfolio_value':d['portfolio_value'],'benchmark_index_value':d.get('benchmark_index_value'),'source_reported_portfolio_return_pct':d.get('portfolio_return_pct'),'source_reported_benchmark_return_pct':d.get('benchmark_return_pct')}):continue
        db.add(PerformanceSnapshot(organization_id=user.organization_id,investor_id=d['investor_id'],source_row_id=row.id,
                                   snapshot_date=date.fromisoformat(d['snapshot_date']),portfolio_value=d['portfolio_value'],
                                   benchmark_index_value=d.get('benchmark_index_value'),
                                   source_reported_portfolio_return_pct=d.get('portfolio_return_pct'),
                                   source_reported_benchmark_return_pct=d.get('benchmark_return_pct'),
                                   return_verification_status='NOT_ASSESSABLE'))
    batch.status='COMMITTED'
    from app.models.entities import utcnow
    batch.committed_at=utcnow()
    db.add(AuditEvent(organization_id=user.organization_id,actor_id=user.id,action='IMPORT_COMMITTED',
                      entity_type='ImportBatch',entity_id=batch.id,detail=json.dumps(batch.sheet_counts)))
    try:
        db.commit()
    except IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=409,detail='Commit conflicts with existing records; no partial import occurred') from e
    return {'id':batch.id,'status':'COMMITTED','committed_counts':batch.sheet_counts,'unchanged_observations_skipped':skipped,'existing_clients_retained':len(prior),
            'warning_count':batch.warning_count,'message':'Import committed. Source warnings remain available for review.'}


def batch_public(batch:ImportBatch)->dict:
    return {'id':batch.id,'filename':batch.filename,'status':batch.status,'created_at':batch.created_at.isoformat(),
            'committed_at':batch.committed_at.isoformat() if batch.committed_at else None,
            'sheet_counts':batch.sheet_counts,'total_rows':batch.total_rows,
            'error_count':batch.error_count,'warning_count':batch.warning_count}
